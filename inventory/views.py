import csv
import os
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db.models import Count, Q
from django.http import FileResponse, HttpResponse, HttpResponseBadRequest, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils.text import slugify
from django.views.decorators.http import require_POST

from .access import can_manage_columns, is_engineer, visible_assets
from .forms import AssetForm
from .grid import (
    ASSET_COLUMNS,
    allowed_column_keys,
    asset_column_value,
    columns_for_keys,
    default_user_column_keys,
    export_column_keys,
    normalized_column_keys,
    user_column_keys,
    user_columns,
)
from .models import (
    Asset,
    AssetAttachment,
    AssetType,
    AssetTypeStatus,
    GridColumnSettings,
    TYPE_MODELS,
    UserColumnPreference,
    all_status_options,
    resolve_asset,
    status_options_for,
)


def _can(user, perm):
    return user.is_superuser or user.has_perm(f"inventory.{perm}")


def _can_change(user, asset):
    """Edit access follows the group's per-type change permission."""
    if user.is_superuser:
        return True
    type_name = asset.asset_type.name.lower() if asset.asset_type_id else ""
    if type_name and user.has_perm(f"inventory.change_{type_name}"):
        return True
    return user.has_perm("inventory.change_asset")


def _can_delete(user, asset):
    """Delete access follows the group's per-type delete permission."""
    if user.is_superuser:
        return True
    type_name = asset.asset_type.name.lower() if asset.asset_type_id else ""
    if type_name and user.has_perm(f"inventory.delete_{type_name}"):
        return True
    return user.has_perm("inventory.delete_asset")


def _next_url(request):
    """Return a safe local URL to use after a column form submission."""
    target = request.POST.get("next") or reverse("inventory:asset_list")
    if url_has_allowed_host_and_scheme(
        target,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return target
    return reverse("inventory:asset_list")


def _filtered_assets(request):
    """Rows shared by the on-screen grid and the export download."""
    assets = visible_assets(request.user).select_related(
        "vendor", "custodian", "asset_type"
    )
    query = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    raw_type = request.GET.get("type", "").strip()
    selected_type = AssetType.objects.filter(pk=raw_type).first() if raw_type else None

    if query:
        assets = assets.filter(
            Q(name__icontains=query)
            | Q(description__icontains=query)
            | Q(custodian__username__icontains=query)
        )
    if status:
        assets = assets.filter(status=status)
    if selected_type:
        assets = assets.filter(asset_type=selected_type)
    return assets, selected_type, query, status


@login_required
def dashboard(request):
    assets = visible_assets(request.user)
    show_all = request.user.is_superuser or is_engineer(request.user)
    counts = {
        row["status"]: row["total"]
        for row in assets.values("status").annotate(total=Count("id"))
    }
    total = assets.count()
    labels = dict(Asset.Status.choices)
    for row in AssetTypeStatus.objects.all():
        labels.setdefault(row.key, row.label)
    keys = list(dict.fromkeys([key for key, _ in Asset.Status.choices] + list(counts)))
    status_cards = [
        {
            "slug": key.replace("_", "-"),
            "label": labels.get(key, key),
            "count": counts.get(key, 0),
        }
        for key in keys
    ]
    recent = assets.select_related("vendor", "custodian", "asset_type")[:6]
    return render(
        request,
        "inventory/dashboard.html",
        {
            "total": total,
            "status_cards": status_cards,
            "recent": recent,
            "active_nav": "dashboard",
            "show_all": show_all,
        },
    )


@login_required
def asset_list(request):
    assets, selected_type, query, status = _filtered_assets(request)
    show_all = request.user.is_superuser

    types = list(AssetType.objects.filter(name__in=TYPE_MODELS))
    counts = {
        row["asset_type__name"]: row["total"]
        for row in visible_assets(request.user)
        .values("asset_type__name")
        .annotate(total=Count("id"))
    }

    # Non-admins don't get the combined "All" view; land them on a type.
    if not show_all and selected_type is None:
        default_type = next(
            (t for t in types if counts.get(t.name)), types[0] if types else None
        )
        if default_type is not None:
            params = request.GET.copy()
            params["type"] = default_type.pk
            return redirect(f"{reverse('inventory:asset_list')}?{params.urlencode()}")

    assets = list(assets)
    for asset in assets:
        asset.can_edit = _can_change(request.user, asset)
        asset.can_delete = _can_delete(request.user, asset)

    total_all = visible_assets(request.user).count()

    allowed_keys = allowed_column_keys()
    visible_keys = user_column_keys(request.user)

    export_params = request.GET.copy()
    export_params.pop("format", None)
    export_params.pop("cols", None)
    export_query = export_params.urlencode()

    return render(
        request,
        "inventory/asset_list.html",
        {
            "assets": assets,
            "statuses": (
                status_options_for(selected_type)
                if selected_type
                else all_status_options()
            ),
            "types": types,
            "type_counts": counts,
            "total_all": total_all,
            "selected_type": selected_type,
            "query": query,
            "selected_status": status,
            "total_shown": len(assets),
            "active_nav": "assets",
            "show_all": show_all,
            "columns": user_columns(request.user),
            "all_columns": ASSET_COLUMNS,
            "allowed_keys": allowed_keys,
            "allowed_count": len(allowed_keys),
            "visible_keys": visible_keys,
            "visible_count": len(visible_keys),
            "default_all_columns": (
                request.user.is_superuser or is_engineer(request.user)
            ),
            "can_manage_columns": can_manage_columns(request.user),
            "export_query": export_query,
        },
    )


@login_required
@require_POST
def columns_save(request):
    """Save (or reset) the signed-in user's personal visible columns."""
    if request.POST.get("reset"):
        UserColumnPreference.objects.filter(user=request.user).delete()
        keys = default_user_column_keys(request.user)
    else:
        allowed = set(allowed_column_keys())
        keys = [
            key
            for key in normalized_column_keys(request.POST.getlist("columns"))
            if key in allowed
        ]
        if not keys:
            error = "Keep at least one column visible."
            if request.headers.get("x-requested-with") == "XMLHttpRequest":
                return JsonResponse({"ok": False, "error": error}, status=400)
            messages.error(request, error)
            return redirect(_next_url(request))

        # Save visible columns and their order
        order = request.POST.getlist("order")
        normalized_order = [
            key for key in normalized_column_keys(order) if key in allowed
        ]
        # Ensure all visible keys are in order (append any missing)
        for key in keys:
            if key not in normalized_order:
                normalized_order.append(key)

        preference, _ = UserColumnPreference.objects.get_or_create(
            user=request.user
        )
        preference.visible_columns = keys
        preference.column_order = normalized_order
        preference.save(update_fields=["visible_columns", "column_order", "updated_at"])

    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return JsonResponse({"ok": True, "columns": keys})

    messages.success(request, "Column view updated.")
    return redirect(_next_url(request))


@login_required
@require_POST
def columns_manage(request):
    """Admin/Engineers update the columns available to every user."""
    if not can_manage_columns(request.user):
        raise PermissionDenied

    keys = normalized_column_keys(request.POST.getlist("columns"))
    if not keys:
        messages.error(request, "Select at least one available column.")
        return redirect(_next_url(request))

    settings_row = GridColumnSettings.load()
    settings_row.allowed_columns = keys
    settings_row.updated_by = request.user
    settings_row.save(
        update_fields=["allowed_columns", "updated_by", "updated_at"]
    )
    messages.success(request, "Available columns updated.")
    return redirect(_next_url(request))


@login_required
def asset_export(request):
    fmt = request.GET.get("format", "").lower()
    if fmt not in {"xlsx", "csv"}:
        return HttpResponseBadRequest("Unsupported export format.")

    assets, selected_type, _, _ = _filtered_assets(request)
    columns = columns_for_keys(export_column_keys(request))
    headers = [column["label"] for column in columns]
    rows = [
        [asset_column_value(asset, column["key"]) for column in columns]
        for asset in assets
    ]

    parts = ["assets"]
    if selected_type:
        parts.append(slugify(selected_type.name) or "type")
    parts.append(timezone.localdate().isoformat())
    filename = "-".join(parts) + f".{fmt}"

    if fmt == "csv":
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response.write("\ufeff")
        writer = csv.writer(response, lineterminator="\r\n")
        writer.writerow(headers)
        writer.writerows(rows)
        return response

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Assets"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        values = [header] + [row[index - 1] for row in rows]
        width = max(len(str(value)) for value in values)
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(width + 2, 12), 42
        )

    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def asset_detail(request, pk):
    asset = resolve_asset(
        get_object_or_404(
            visible_assets(request.user).select_related(
                "vendor", "custodian", "created_by", "updated_by", "asset_type"
            ),
            pk=pk,
        )
    )
    return render(
        request,
        "inventory/asset_detail.html",
        {
            "asset": asset,
            "attachments": asset.attachments.select_related("uploaded_by").all(),
            "can_change": _can_change(request.user, asset),
            "can_delete": _can_delete(request.user, asset),
            "active_nav": "assets",
        },
    )


ATTACHMENT_MAX_BYTES = 20 * 1024 * 1024
ATTACHMENT_MAX_FILES = 10
ATTACHMENT_ALLOWED_EXTENSIONS = {
    ".pdf", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".txt", ".csv",
    ".xlsx", ".docx", ".doc", ".zip", ".log", ".cfg", ".conf", ".json", ".xml",
}


@login_required
@require_POST
def asset_attachment_upload(request, pk):
    asset = get_object_or_404(visible_assets(request.user), pk=pk)
    if not _can_change(request.user, asset):
        raise PermissionDenied

    uploads = request.FILES.getlist("files")
    if not uploads:
        messages.error(request, "Choose at least one file to upload.")
        return redirect(asset)
    if len(uploads) > ATTACHMENT_MAX_FILES:
        messages.error(request, f"Upload no more than {ATTACHMENT_MAX_FILES} files at a time.")
        return redirect(asset)

    errors = []
    for uploaded in uploads:
        name = os.path.basename(uploaded.name).replace("\\", "_")[:255]
        extension = os.path.splitext(name)[1].lower()
        if extension not in ATTACHMENT_ALLOWED_EXTENSIONS:
            errors.append(f"{name}: this file type is not allowed.")
        elif uploaded.size > ATTACHMENT_MAX_BYTES:
            errors.append(f"{name}: exceeds the 20 MB per-file limit.")
    if errors:
        for error in errors:
            messages.error(request, error)
        return redirect(asset)

    for uploaded in uploads:
        name = os.path.basename(uploaded.name).replace("\\", "_")[:255]
        AssetAttachment.objects.create(
            asset=asset,
            file=uploaded,
            original_name=name,
            file_size=uploaded.size,
            uploaded_by=request.user,
        )
    messages.success(request, f"{len(uploads)} attachment(s) uploaded.")
    return redirect(asset)


@login_required
def asset_attachment_download(request, pk, attachment_pk):
    asset = get_object_or_404(visible_assets(request.user), pk=pk)
    attachment = get_object_or_404(AssetAttachment, pk=attachment_pk, asset=asset)
    try:
        file_handle = attachment.file.open("rb")
    except (FileNotFoundError, OSError):
        messages.error(request, "The stored attachment file is unavailable.")
        return redirect(asset)
    response = FileResponse(
        file_handle, as_attachment=True, filename=attachment.original_name
    )
    response["X-Content-Type-Options"] = "nosniff"
    response["Cache-Control"] = "private, no-store"
    return response


@login_required
@require_POST
def asset_attachment_delete(request, pk, attachment_pk):
    asset = get_object_or_404(visible_assets(request.user), pk=pk)
    if not _can_change(request.user, asset):
        raise PermissionDenied
    attachment = get_object_or_404(AssetAttachment, pk=attachment_pk, asset=asset)
    attachment.delete()
    messages.success(request, "Attachment deleted.")
    return redirect(asset)


@login_required
def asset_create(request):
    if not _can(request.user, "add_asset"):
        raise PermissionDenied
    initial = {}
    raw_type = request.GET.get("type", "").strip()
    if raw_type:
        initial["asset_type"] = raw_type
    if request.method == "POST":
        form = AssetForm(request.POST, request.FILES, user=request.user)
        if form.is_valid():
            asset = form.save(commit=False)
            asset.created_by = request.user
            asset.save()
            messages.success(request, f"Asset “{asset.name}” created.")
            return redirect(asset)
    else:
        form = AssetForm(user=request.user, initial=initial or None)
    return render(
        request,
        "inventory/asset_form.html",
        {"form": form, "active_nav": "assets", "title": "New asset"},
    )


@login_required
def asset_edit(request, pk):
    asset = resolve_asset(
        get_object_or_404(visible_assets(request.user), pk=pk)
    )
    if not _can_change(request.user, asset):
        raise PermissionDenied
    if request.method == "POST":
        form = AssetForm(request.POST, instance=asset, user=request.user)
        if form.is_valid():
            asset = form.save(commit=False)
            asset.updated_by = request.user
            asset.save()
            messages.success(request, f"Asset “{asset.name}” updated.")
            return redirect(asset)
    else:
        form = AssetForm(instance=asset, user=request.user)
    return render(
        request,
        "inventory/asset_form.html",
        {"form": form, "asset": asset, "active_nav": "assets", "title": "Edit asset"},
    )


@login_required
def asset_delete(request, pk):
    asset = resolve_asset(
        get_object_or_404(visible_assets(request.user), pk=pk)
    )
    if not _can_delete(request.user, asset):
        raise PermissionDenied
    if request.method == "POST":
        name = asset.name
        asset.delete()
        messages.success(request, f"Asset “{name}” deleted.")
        return redirect("inventory:asset_list")
    return render(
        request,
        "inventory/asset_confirm_delete.html",
        {"asset": asset, "active_nav": "assets"},
    )